from __future__ import annotations

import csv
import re
from datetime import date, datetime, time, timezone
from io import BytesIO, StringIO
from typing import Optional
from uuid import UUID, uuid4
from zoneinfo import ZoneInfo

from sqlalchemy import delete, func, inspect, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from sqlalchemy.orm.attributes import set_committed_value

from app.models.business import Business, Location, Role
from app.models.common import AssignmentStatus, ShiftLifecycleStatus
from app.models.workforce import (
    Employee,
    EmployeeAvailabilityRule,
    EmployeeLocation,
    EmployeeRole,
    EmployeeWorkPermit,
)
from app.schemas.workforce import (
    EmployeeAvailabilityRuleCreate,
    EmployeeAvailabilityRuleReplace,
    EmployeeBulkImportRead,
    EmployeeCreate,
    EmployeeEnrollAtLocationCreate,
    EmployeeEnrollmentRead,
    EmployeeImportErrorRead,
    EmployeeLocationCreate,
    EmployeeNotificationPreferencesUpdate,
    EmployeeLocationUpsert,
    EmployeeRoleCreate,
    EmployeeRoleUpsert,
    EmployeeWorkPermitCreate,
    EmployeeWorkPermitTemplateRead,
    EmployeeUpdate,
)
from app.services import work_permit_rules


_DEFAULT_EMPLOYEE_NOTIFICATION_PREFERENCES = {
    "schedule_publish_email_enabled": True,
    "schedule_publish_sms_enabled": False,
    "email_opted_out_at": None,
    "sms_opted_out_at": None,
    "email_opt_out_reason": None,
    "sms_opt_out_reason": None,
}

EMPLOYEE_IMPORT_HEADERS = (
    "first_name",
    "last_name",
    "phone_number",
    "email_address",
    "date_of_birth",
    "minor_school_status",
    "work_permit_number",
    "work_permit_expires_on",
)

EMPLOYEE_IMPORT_REQUIRED_FIELDS = (
    "full_name",
    "email",
    "phone_e164",
)

EMPLOYEE_IMPORT_HEADER_ALIASES = {
    "full_name": "full_name",
    "employee_name": "full_name",
    "name": "full_name",
    "employee_full_name": "full_name",
    "employee": "full_name",
    "staff_name": "full_name",
    "worker_name": "full_name",
    "preferred_name": "preferred_name",
    "email": "email",
    "email_address": "email",
    "email_addr": "email",
    "emailaddress": "email",
    "e_mail": "email",
    "work_email": "email",
    "personal_email": "email",
    "email_id": "email",
    "phone": "phone_e164",
    "phone_number": "phone_e164",
    "phone_e164": "phone_e164",
    "phone_num": "phone_e164",
    "phonenumber": "phone_e164",
    "telephone": "phone_e164",
    "contact_number": "phone_e164",
    "contact_phone": "phone_e164",
    "employee_phone": "phone_e164",
    "work_phone": "phone_e164",
    "mobile": "phone_e164",
    "mobile_phone": "phone_e164",
    "mobile_number": "phone_e164",
    "mobile_no": "phone_e164",
    "cell": "phone_e164",
    "cell_phone": "phone_e164",
    "cell_number": "phone_e164",
    "employee_number": "employee_number",
    "employee_id": "employee_number",
    "external_ref": "external_ref",
    "external_id": "external_ref",
    "employment_type": "employment_type",
    "employment_status": "employment_type",
    "pay_rate": "base_hourly_rate_cents",
    "hourly_rate": "base_hourly_rate_cents",
    "base_hourly_rate": "base_hourly_rate_cents",
    "regular_rate": "compliance_regular_rate_cents",
    "regular_rate_of_pay": "compliance_regular_rate_cents",
    "compliance_regular_rate": "compliance_regular_rate_cents",
    "compliance_regular_rate_of_pay": "compliance_regular_rate_cents",
    "meal_rest_premium_rate": "compliance_regular_rate_cents",
    "date_of_birth": "date_of_birth",
    "birth_date": "date_of_birth",
    "dob": "date_of_birth",
    "minor_school_status": "minor_school_status",
    "school_status": "minor_school_status",
    "school_calendar_status": "minor_school_status",
    "work_permit_number": "work_permit_number",
    "permit_number": "work_permit_number",
    "work_permit_id": "work_permit_number",
    "permit_id": "work_permit_number",
    "work_permit_expires_on": "work_permit_expires_on",
    "work_permit_expiration_date": "work_permit_expires_on",
    "permit_expires_on": "work_permit_expires_on",
    "permit_expiration_date": "work_permit_expires_on",
    "hire_date": "hire_date",
    "start_date": "hire_date",
    "notes": "notes",
    "first_name": "first_name",
    "firstname": "first_name",
    "first": "first_name",
    "given_name": "first_name",
    "givenname": "first_name",
    "last_name": "last_name",
    "lastname": "last_name",
    "last": "last_name",
    "family_name": "last_name",
    "familyname": "last_name",
    "surname": "last_name",
}

def _require_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
    except ModuleNotFoundError as exc:
        raise RuntimeError("employee_import_xlsx_dependency_missing") from exc
    return Workbook, load_workbook, Font, PatternFill


def list_work_permit_templates() -> list[EmployeeWorkPermitTemplateRead]:
    return [
        EmployeeWorkPermitTemplateRead.model_validate(item)
        for item in work_permit_rules.list_work_permit_templates()
    ]


async def _require_business(session: AsyncSession, business_id: UUID) -> Business:
    business = await session.get(Business, business_id)
    if business is None:
        raise LookupError("business_not_found")
    return business


async def _list_employee_locations(
    session: AsyncSession,
    employee_id: UUID,
) -> list[EmployeeLocation]:
    result = await session.execute(
        select(EmployeeLocation)
        .options(selectinload(EmployeeLocation.location))
        .where(EmployeeLocation.employee_id == employee_id)
        .order_by(EmployeeLocation.created_at.asc())
    )
    return list(result.scalars().all())


async def _list_employee_roles(
    session: AsyncSession,
    employee_id: UUID,
) -> list[EmployeeRole]:
    result = await session.execute(
        select(EmployeeRole)
        .options(selectinload(EmployeeRole.role))
        .where(EmployeeRole.employee_id == employee_id)
        .order_by(EmployeeRole.created_at.asc())
    )
    return list(result.scalars().all())


async def _list_employee_work_permits(
    session: AsyncSession,
    employee_id: UUID,
) -> list[EmployeeWorkPermit]:
    result = await session.execute(
        select(EmployeeWorkPermit)
        .where(EmployeeWorkPermit.employee_id == employee_id)
        .order_by(
            EmployeeWorkPermit.effective_start_date.asc().nullsfirst(),
            EmployeeWorkPermit.created_at.asc(),
        )
    )
    return list(result.scalars().all())


def _employee_query():
    return select(Employee).options(
        selectinload(Employee.employee_roles).selectinload(EmployeeRole.role),
        selectinload(Employee.employee_locations).selectinload(EmployeeLocation.location),
    )


def _work_permit_sort_key(permit: EmployeeWorkPermit | EmployeeWorkPermitCreate) -> tuple[date, date, str]:
    effective_end_date = getattr(permit, "effective_end_date", None) or date.max
    effective_start_date = getattr(permit, "effective_start_date", None) or date.min
    permit_number = str(getattr(permit, "permit_number", "") or "")
    return effective_end_date, effective_start_date, permit_number


def _serialized_work_permit_rule_profile(
    value: object | None,
) -> dict[str, object] | None:
    return work_permit_rules.resolve_work_permit_rule_profile_payload(value)


def _work_permit_rule_profile_from_metadata(
    metadata: object | None,
) -> dict[str, object] | None:
    if not isinstance(metadata, dict):
        return None
    return _serialized_work_permit_rule_profile(metadata.get("rule_profile"))


def _employee_work_permit_rule_profile_snapshot(
    employee: Employee,
) -> dict[str, object] | None:
    if not isinstance(employee.employee_metadata, dict):
        return None
    return _serialized_work_permit_rule_profile(
        employee.employee_metadata.get("work_permit_rule_profile")
    )


def _set_employee_work_permit_rule_profile_snapshot(
    employee: Employee,
    rule_profile: dict[str, object] | None,
) -> None:
    metadata = dict(employee.employee_metadata or {})
    if rule_profile:
        metadata["work_permit_rule_profile"] = rule_profile
    else:
        metadata.pop("work_permit_rule_profile", None)
    employee.employee_metadata = metadata


def _hydrate_employee_defaults(employee: Employee) -> None:
    now = datetime.now(timezone.utc)
    if employee.id is None:
        employee.id = uuid4()
    if employee.created_at is None:
        employee.created_at = now
    if employee.updated_at is None:
        employee.updated_at = now


def _hydrate_work_permit_defaults(employee: Employee, permit: EmployeeWorkPermit) -> None:
    _hydrate_employee_defaults(employee)
    now = datetime.now(timezone.utc)
    permit.employee_id = employee.id
    if permit.id is None:
        permit.id = uuid4()
    if permit.created_at is None:
        permit.created_at = now
    if permit.updated_at is None:
        permit.updated_at = now


def _sync_employee_work_permit_snapshot(
    employee: Employee,
    permits: list[EmployeeWorkPermit] | list[EmployeeWorkPermitCreate],
) -> None:
    if not permits:
        employee.work_permit_number = None
        employee.work_permit_effective_start_on = None
        employee.work_permit_expires_on = None
        employee.work_permit_max_daily_minutes = None
        employee.work_permit_max_weekly_minutes = None
        employee.work_permit_earliest_start_local_time = None
        employee.work_permit_latest_end_local_time = None
        _set_employee_work_permit_rule_profile_snapshot(employee, None)
        return
    best = sorted(permits, key=_work_permit_sort_key, reverse=True)[0]
    employee.work_permit_number = str(getattr(best, "permit_number", "") or "").strip() or None
    employee.work_permit_effective_start_on = getattr(best, "effective_start_date", None)
    employee.work_permit_expires_on = getattr(best, "effective_end_date", None)
    employee.work_permit_max_daily_minutes = getattr(best, "max_daily_minutes", None)
    employee.work_permit_max_weekly_minutes = getattr(best, "max_weekly_minutes", None)
    employee.work_permit_earliest_start_local_time = getattr(best, "earliest_start_local_time", None)
    employee.work_permit_latest_end_local_time = getattr(best, "latest_end_local_time", None)
    _set_employee_work_permit_rule_profile_snapshot(
        employee,
        _serialized_work_permit_rule_profile(
            getattr(best, "rule_profile", None)
            or _work_permit_rule_profile_from_metadata(
                getattr(best, "permit_metadata", None)
            )
        ),
    )


def _loaded_employee_work_permits(employee: Employee) -> list[EmployeeWorkPermit]:
    try:
        state = inspect(employee)
        if "work_permits" in getattr(state, "unloaded", ()):
            return []
    except Exception:
        return []
    return list(getattr(employee, "work_permits", []) or [])


def _permit_covers_date(permit: EmployeeWorkPermit, reference_date: date) -> bool:
    start_on = permit.effective_start_date
    end_on = permit.effective_end_date
    if start_on is not None and start_on > reference_date:
        return False
    if end_on is not None and end_on < reference_date:
        return False
    return True


def _active_work_permit_sort_key(permit: EmployeeWorkPermit) -> tuple[date, date, str]:
    effective_start_date = permit.effective_start_date or date.min
    effective_end_date = permit.effective_end_date or date.max
    permit_number = str(permit.permit_number or "")
    return effective_start_date, effective_end_date, permit_number


def _snapshot_work_permit_context(employee: Employee, *, reference_date: date) -> dict[str, object]:
    permit_number = str(getattr(employee, "work_permit_number", "") or "").strip() or None
    effective_start_on = getattr(employee, "work_permit_effective_start_on", None)
    expires_on = getattr(employee, "work_permit_expires_on", None)
    is_active = bool(permit_number)
    if effective_start_on is not None and effective_start_on > reference_date:
        is_active = False
    if expires_on is not None and expires_on < reference_date:
        is_active = False
    return {
        "permit_number": permit_number,
        "effective_start_on": effective_start_on,
        "expires_on": expires_on,
        "max_daily_minutes": getattr(employee, "work_permit_max_daily_minutes", None),
        "max_weekly_minutes": getattr(employee, "work_permit_max_weekly_minutes", None),
        "earliest_start_local_time": getattr(employee, "work_permit_earliest_start_local_time", None),
        "latest_end_local_time": getattr(employee, "work_permit_latest_end_local_time", None),
        "rule_profile": _employee_work_permit_rule_profile_snapshot(employee),
        "is_active": is_active,
        "source": "employee_snapshot",
    }


def resolve_employee_work_permit_context(
    employee: Employee,
    *,
    shift_starts_at: datetime,
    timezone_name: str | None,
) -> dict[str, object]:
    try:
        shift_timezone = ZoneInfo(str(timezone_name or "").strip() or "UTC")
    except Exception:
        shift_timezone = ZoneInfo("UTC")
    shift_local_date = shift_starts_at.astimezone(shift_timezone).date()
    loaded_permits = _loaded_employee_work_permits(employee)
    if not loaded_permits:
        return _snapshot_work_permit_context(employee, reference_date=shift_local_date)

    active_permits = [permit for permit in loaded_permits if _permit_covers_date(permit, shift_local_date)]
    if active_permits:
        permit = sorted(active_permits, key=_active_work_permit_sort_key, reverse=True)[0]
        return {
            "permit_number": permit.permit_number,
            "effective_start_on": permit.effective_start_date,
            "expires_on": permit.effective_end_date,
            "max_daily_minutes": permit.max_daily_minutes,
            "max_weekly_minutes": permit.max_weekly_minutes,
            "earliest_start_local_time": permit.earliest_start_local_time,
            "latest_end_local_time": permit.latest_end_local_time,
            "rule_profile": _serialized_work_permit_rule_profile(permit.rule_profile),
            "is_active": True,
            "source": "employee_work_permit",
        }

    future_permits = sorted(
        [
            permit
            for permit in loaded_permits
            if permit.effective_start_date is not None and permit.effective_start_date > shift_local_date
        ],
        key=lambda permit: (
            permit.effective_start_date or date.max,
            permit.effective_end_date or date.max,
            permit.permit_number,
        ),
    )
    if future_permits:
        permit = future_permits[0]
    else:
        past_permits = sorted(
            [
                permit
                for permit in loaded_permits
                if permit.effective_end_date is not None and permit.effective_end_date < shift_local_date
            ],
            key=lambda permit: (
                permit.effective_end_date or date.min,
                permit.effective_start_date or date.min,
                permit.permit_number,
            ),
            reverse=True,
        )
        permit = past_permits[0] if past_permits else sorted(loaded_permits, key=_work_permit_sort_key, reverse=True)[0]

    return {
        "permit_number": permit.permit_number,
        "effective_start_on": permit.effective_start_date,
        "expires_on": permit.effective_end_date,
        "max_daily_minutes": permit.max_daily_minutes,
        "max_weekly_minutes": permit.max_weekly_minutes,
        "earliest_start_local_time": permit.earliest_start_local_time,
        "latest_end_local_time": permit.latest_end_local_time,
        "rule_profile": _serialized_work_permit_rule_profile(permit.rule_profile),
        "is_active": False,
        "source": "employee_work_permit",
    }


def _normalize_employee_phone(value: str | None) -> str | None:
    normalized = str(value or "").strip()
    if not normalized:
        return None
    digits = re.sub(r"\D", "", normalized)
    if normalized.startswith("+") and 10 <= len(digits) <= 15:
        return f"+{digits}"
    if len(digits) == 10:
        return f"+1{digits}"
    if len(digits) == 11 and digits.startswith("1"):
        return f"+{digits}"
    return normalized


def _normalize_employee_email(value: str | None) -> str | None:
    normalized = str(value or "").strip().lower()
    return normalized or None


def _normalize_employee_external_ref(value: str | None) -> str | None:
    normalized = str(value or "").strip()
    return normalized or None


def _normalize_minor_school_status(value: object | None) -> str | None:
    normalized = str(value or "").strip().lower()
    if not normalized:
        return None
    aliases = {
        "unknown": "unknown",
        "in_session": "in_session",
        "school_year": "in_session",
        "school_in_session": "in_session",
        "on_break": "summer_break",
        "summer_break": "summer_break",
        "summer": "summer_break",
        "not_enrolled": "not_enrolled",
    }
    resolved = aliases.get(normalized)
    if resolved is None:
        raise ValueError("invalid_minor_school_status")
    return resolved


def _normalize_notification_datetime(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return None


def normalized_employee_notification_preferences(metadata: dict | None) -> dict[str, object]:
    raw = dict(((metadata or {}).get("notification_preferences")) or {})
    return {
        "schedule_publish_email_enabled": bool(
            raw.get(
                "schedule_publish_email_enabled",
                _DEFAULT_EMPLOYEE_NOTIFICATION_PREFERENCES["schedule_publish_email_enabled"],
            )
        ),
        "schedule_publish_sms_enabled": bool(
            raw.get(
                "schedule_publish_sms_enabled",
                _DEFAULT_EMPLOYEE_NOTIFICATION_PREFERENCES["schedule_publish_sms_enabled"],
            )
        ),
        "email_opted_out_at": _normalize_notification_datetime(raw.get("email_opted_out_at")),
        "sms_opted_out_at": _normalize_notification_datetime(raw.get("sms_opted_out_at")),
        "email_opt_out_reason": str(raw.get("email_opt_out_reason") or "").strip() or None,
        "sms_opt_out_reason": str(raw.get("sms_opt_out_reason") or "").strip() or None,
    }


def _serialized_employee_notification_preferences(preferences: dict[str, object]) -> dict[str, object]:
    return {
        "schedule_publish_email_enabled": bool(preferences["schedule_publish_email_enabled"]),
        "schedule_publish_sms_enabled": bool(preferences["schedule_publish_sms_enabled"]),
        "email_opted_out_at": (
            preferences["email_opted_out_at"].isoformat()
            if isinstance(preferences.get("email_opted_out_at"), datetime)
            else None
        ),
        "sms_opted_out_at": (
            preferences["sms_opted_out_at"].isoformat()
            if isinstance(preferences.get("sms_opted_out_at"), datetime)
            else None
        ),
        "email_opt_out_reason": str(preferences.get("email_opt_out_reason") or "").strip() or None,
        "sms_opt_out_reason": str(preferences.get("sms_opt_out_reason") or "").strip() or None,
    }


def _apply_employee_notification_preferences_update(
    employee: Employee,
    payload: EmployeeNotificationPreferencesUpdate,
) -> None:
    preferences = normalized_employee_notification_preferences(employee.employee_metadata)
    for field_name in payload.model_fields_set:
        preferences[field_name] = getattr(payload, field_name)
    employee.employee_metadata = {
        **(employee.employee_metadata or {}),
        "notification_preferences": _serialized_employee_notification_preferences(preferences),
    }


def _attach_employee_notification_preferences(employee: Employee) -> Employee:
    setattr(
        employee,
        "notification_preferences",
        normalized_employee_notification_preferences(employee.employee_metadata),
    )
    return employee


async def _find_duplicate_employee(
    session: AsyncSession,
    business_id: UUID,
    *,
    external_ref: str | None,
    phone_e164: str | None,
    email: str | None,
    exclude_employee_id: UUID | None = None,
) -> tuple[str, Employee] | None:
    normalized_external_ref = _normalize_employee_external_ref(external_ref)
    normalized_phone = _normalize_employee_phone(phone_e164)
    normalized_email = _normalize_employee_email(email)

    stmt = select(Employee).where(Employee.business_id == business_id)
    if exclude_employee_id is not None:
        stmt = stmt.where(Employee.id != exclude_employee_id)
    result = await session.execute(stmt)
    employees = list(result.scalars().all())

    if normalized_external_ref:
        for employee in employees:
            if _normalize_employee_external_ref(employee.external_ref) == normalized_external_ref:
                return ("external_ref", employee)

    if normalized_phone:
        for employee in employees:
            if _normalize_employee_phone(employee.phone_e164) == normalized_phone:
                return ("phone_e164", employee)

    if normalized_email:
        for employee in employees:
            if _normalize_employee_email(employee.email) == normalized_email:
                return ("email", employee)

    return None


EMPLOYEE_DELETE_BLOCKING_ASSIGNMENT_STATUSES = (
    AssignmentStatus.assigned,
    AssignmentStatus.accepted,
    AssignmentStatus.completed,
)


async def _list_employee_availability_rules(
    session: AsyncSession,
    employee_id: UUID,
) -> list[EmployeeAvailabilityRule]:
    result = await session.execute(
        select(EmployeeAvailabilityRule)
        .where(EmployeeAvailabilityRule.employee_id == employee_id)
        .order_by(
            EmployeeAvailabilityRule.day_of_week.asc(),
            EmployeeAvailabilityRule.start_local_time.asc(),
            EmployeeAvailabilityRule.priority.asc(),
            EmployeeAvailabilityRule.created_at.asc(),
        )
    )
    return list(result.scalars().all())


def _self_service_employee_name(
    *,
    full_name: str | None,
    email: str | None,
    phone_e164: str | None,
) -> str:
    normalized_name = full_name.strip() if full_name else ""
    if normalized_name:
        return normalized_name

    normalized_email = email.strip() if email else ""
    if normalized_email:
        return normalized_email

    normalized_phone = phone_e164.strip() if phone_e164 else ""
    if normalized_phone:
        return normalized_phone

    return "Backfill User"


def _default_availability_timezone(
    *,
    business: Business,
    primary_location: Location | None,
) -> str:
    location_timezone = str(getattr(primary_location, "timezone", "") or "").strip()
    if location_timezone:
        return location_timezone
    business_timezone = str(getattr(business, "timezone", "") or "").strip()
    return business_timezone or "UTC"


def _default_weekly_availability_rules(
    *,
    employee_id: UUID,
    timezone_name: str,
    source: str,
) -> list[EmployeeAvailabilityRule]:
    return [
        EmployeeAvailabilityRule(
            employee_id=employee_id,
            day_of_week=day_of_week,
            start_local_time=time(0, 0),
            end_local_time=time(23, 59),
            timezone=timezone_name,
            availability_type="available",
            priority=0,
            availability_metadata={"source": source, "preset": "all_days"},
        )
        for day_of_week in range(7)
    ]


def _pick_linkable_employee_match(
    matches: list[Employee],
    *,
    user_id: UUID | None,
) -> Employee | None:
    if user_id is None:
        return matches[0] if len(matches) == 1 else None

    eligible = [
        employee
        for employee in matches
        if employee.user_id is None or employee.user_id == user_id
    ]
    return eligible[0] if len(eligible) == 1 else None


async def _find_self_service_employee(
    session: AsyncSession,
    business_id: UUID,
    *,
    user_id: UUID | None,
    email: str | None,
    phone_e164: str | None,
    full_name: str | None,
) -> Employee | None:
    if user_id is not None:
        linked_employee = await session.scalar(
            _employee_query().where(
                Employee.business_id == business_id,
                Employee.user_id == user_id,
            )
        )
        if linked_employee is not None:
            return linked_employee

    normalized_phone = phone_e164.strip() if phone_e164 else ""
    if normalized_phone:
        phone_result = await session.execute(
            _employee_query().where(
                Employee.business_id == business_id,
                Employee.phone_e164 == normalized_phone,
            )
        )
        phone_match = _pick_linkable_employee_match(
            list(phone_result.scalars().all()),
            user_id=user_id,
        )
        if phone_match is not None:
            return phone_match

    normalized_email = email.strip().lower() if email else ""
    if normalized_email:
        email_result = await session.execute(
            _employee_query().where(
                Employee.business_id == business_id,
                func.lower(Employee.email) == normalized_email,
            )
        )
        email_match = _pick_linkable_employee_match(
            list(email_result.scalars().all()),
            user_id=user_id,
        )
        if email_match is not None:
            return email_match

    normalized_name = full_name.strip() if full_name else ""
    if normalized_name:
        result = await session.execute(
            _employee_query().where(
                Employee.business_id == business_id,
                Employee.full_name == normalized_name,
            )
        )
        name_match = _pick_linkable_employee_match(
            list(result.scalars().all()),
            user_id=user_id,
        )
        if name_match is not None:
            return name_match

    return None


async def _resolve_self_service_employee(
    session: AsyncSession,
    business_id: UUID,
    *,
    user_id: UUID | None,
    email: str | None,
    phone_e164: str | None,
    full_name: str | None,
) -> Employee:
    employee = await _find_self_service_employee(
        session,
        business_id,
        user_id=user_id,
        email=email,
        phone_e164=phone_e164,
        full_name=full_name,
    )
    if employee is not None:
        return employee

    raise LookupError("employee_self_not_found")


async def _link_employee_to_user(
    session: AsyncSession,
    employee: Employee,
    *,
    user_id: UUID | None,
) -> Employee:
    if user_id is None or employee.user_id == user_id:
        return employee
    if employee.user_id is not None:
        raise LookupError("employee_self_link_conflict")
    employee.user_id = user_id
    await session.flush()
    return employee


async def _get_or_create_self_service_employee(
    session: AsyncSession,
    business_id: UUID,
    *,
    user_id: UUID | None,
    email: str | None,
    phone_e164: str | None,
    full_name: str | None,
) -> Employee:
    employee = await _find_self_service_employee(
        session,
        business_id,
        user_id=user_id,
        email=email,
        phone_e164=phone_e164,
        full_name=full_name,
    )
    if employee is not None:
        return await _link_employee_to_user(session, employee, user_id=user_id)

    location_result = await session.execute(
        select(Location.id)
        .where(Location.business_id == business_id, Location.is_active.is_(True))
        .order_by(Location.created_at.asc())
    )
    active_location_ids = list(location_result.scalars().all())
    primary_location_id = active_location_ids[0] if len(active_location_ids) == 1 else None

    employee = await create_employee(
        session,
        business_id,
        EmployeeCreate(
            full_name=_self_service_employee_name(
                full_name=full_name,
                email=email,
                phone_e164=phone_e164,
            ),
            phone_e164=phone_e164.strip() if phone_e164 else None,
            email=email.strip().lower() if email else None,
            primary_location_id=primary_location_id,
            employee_metadata={
                "source": "self_service_availability",
                "auto_created_from_user": True,
            },
        ),
        linked_user_id=user_id,
    )
    return employee


def _normalize_import_header(value: object) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower())
    return normalized.strip("_")


def _normalize_import_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _parse_hire_date(raw: object) -> date | None:
    return _parse_optional_date(raw, error_code="invalid_hire_date")


def _parse_date_of_birth(raw: object) -> date | None:
    return _parse_optional_date(raw, error_code="invalid_date_of_birth")


def _parse_work_permit_expires_on(raw: object) -> date | None:
    return _parse_optional_date(raw, error_code="invalid_work_permit_expires_on")


def _parse_optional_date(raw: object, *, error_code: str) -> date | None:
    if raw in (None, ""):
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw

    normalized = str(raw).strip()
    if not normalized:
        return None

    for candidate in (normalized, normalized.replace("/", "-")):
        try:
            return date.fromisoformat(candidate)
        except ValueError:
            pass

    for fmt in ("%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y", "%m-%d-%y"):
        try:
            return datetime.strptime(normalized, fmt).date()
        except ValueError:
            continue

    raise ValueError(error_code)


def _parse_hourly_rate_cents(raw: object, *, error_code: str) -> int | None:
    if raw in (None, ""):
        return None
    if isinstance(raw, int):
        return max(0, raw)

    normalized = str(raw).strip()
    if not normalized:
        return None
    normalized = normalized.replace("$", "").replace(",", "")
    try:
        return max(0, int(round(float(normalized) * 100)))
    except ValueError as exc:
        raise ValueError(error_code) from exc


def _parse_base_hourly_rate_cents(raw: object) -> int | None:
    return _parse_hourly_rate_cents(raw, error_code="invalid_base_hourly_rate")


def _parse_compliance_regular_rate_cents(raw: object) -> int | None:
    return _parse_hourly_rate_cents(raw, error_code="invalid_compliance_regular_rate")


def _canonicalize_import_row(raw_row: dict[object, object]) -> dict[str, object]:
    canonical: dict[str, object] = {}
    first_name = ""
    last_name = ""

    for header, value in raw_row.items():
        normalized_header = _normalize_import_header(header)
        if not normalized_header:
            continue
        target = EMPLOYEE_IMPORT_HEADER_ALIASES.get(normalized_header)
        if target is None:
            continue
        if target == "first_name":
            first_name = _normalize_import_value(value)
            continue
        if target == "last_name":
            last_name = _normalize_import_value(value)
            continue
        canonical[target] = _normalize_import_value(value)

    if not canonical.get("full_name"):
        full_name = " ".join(part for part in (first_name, last_name) if part)
        if full_name:
            canonical["full_name"] = full_name

    return canonical


def parse_employee_import_file(
    filename: str,
    content: bytes,
    *,
    default_location_id: UUID | None = None,
) -> tuple[list[EmployeeCreate], list[EmployeeImportErrorRead]]:
    normalized_name = filename.lower().strip()
    if normalized_name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(StringIO(text))
        rows = [(index, row) for index, row in enumerate(reader, start=2)]
    elif normalized_name.endswith(".xlsx"):
        _, load_workbook, _, _ = _require_openpyxl()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if headers is None:
            return [], []
        rows = []
        for index, values in enumerate(iterator, start=2):
            row = {
                headers[column_index]: values[column_index]
                for column_index in range(len(headers))
            }
            rows.append((index, row))
    else:
        raise ValueError("employee_import_file_type_unsupported")

    employees: list[EmployeeCreate] = []
    errors: list[EmployeeImportErrorRead] = []

    for row_number, raw_row in rows:
        canonical = _canonicalize_import_row(raw_row)
        if not any(str(value).strip() for value in canonical.values() if value is not None):
            continue

        missing_fields = [
            field_name
            for field_name in EMPLOYEE_IMPORT_REQUIRED_FIELDS
            if not str(canonical.get(field_name) or "").strip()
        ]
        if missing_fields:
            errors.append(
                EmployeeImportErrorRead(
                    row_number=row_number,
                    message=f"missing_required_fields:{','.join(missing_fields)}",
                )
            )
            continue

        try:
            normalized_phone = _normalize_employee_phone(
                str(canonical.get("phone_e164") or "").strip() or None,
            )
            employee = EmployeeCreate(
                full_name=str(canonical.get("full_name") or "").strip(),
                preferred_name=str(canonical.get("preferred_name") or "").strip() or None,
                email=_normalize_employee_email(
                    str(canonical.get("email") or "").strip() or None,
                ),
                phone_e164=normalized_phone,
                employee_number=str(canonical.get("employee_number") or "").strip() or None,
                external_ref=str(canonical.get("external_ref") or "").strip() or None,
                base_hourly_rate_cents=_parse_base_hourly_rate_cents(
                    canonical.get("base_hourly_rate_cents")
                ),
                compliance_regular_rate_cents=_parse_compliance_regular_rate_cents(
                    canonical.get("compliance_regular_rate_cents")
                ),
                date_of_birth=_parse_date_of_birth(canonical.get("date_of_birth")),
                minor_school_status=_normalize_minor_school_status(canonical.get("minor_school_status")),
                work_permit_number=str(canonical.get("work_permit_number") or "").strip() or None,
                work_permit_expires_on=_parse_work_permit_expires_on(
                    canonical.get("work_permit_expires_on")
                ),
                employment_type=str(canonical.get("employment_type") or "").strip() or None,
                primary_location_id=default_location_id,
                hire_date=_parse_hire_date(canonical.get("hire_date")),
                notes=str(canonical.get("notes") or "").strip() or None,
                employee_metadata={
                    "source": "bulk_import",
                    "source_row_number": row_number,
                },
            )
        except ValueError as exc:
            errors.append(
                EmployeeImportErrorRead(
                    row_number=row_number,
                    message=str(exc),
                )
            )
            continue

        employees.append(employee)

    if not employees and not errors:
        raise ValueError("employee_import_no_rows")

    return employees, errors


def build_employee_import_template() -> bytes:
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(EMPLOYEE_IMPORT_HEADERS)
    writer.writerow(
        [
            "Taylor",
            "Smith",
            "+15555550123",
            "taylor@example.com",
            "1998-04-12",
            "unknown",
            "",
            "",
        ]
    )
    return buffer.getvalue().encode("utf-8")


async def list_employees(session: AsyncSession, business_id: UUID) -> list[Employee]:
    result = await session.execute(
        _employee_query()
        .where(Employee.business_id == business_id)
        .order_by(Employee.created_at.desc())
    )
    return [_attach_employee_notification_preferences(employee) for employee in result.scalars().all()]


async def get_employee(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
) -> Optional[Employee]:
    employee = await session.scalar(
        _employee_query().where(
            Employee.id == employee_id,
            Employee.business_id == business_id,
        )
    )
    if employee is None:
        return None
    return _attach_employee_notification_preferences(employee)


async def _require_employee(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
) -> Employee:
    employee = await get_employee(session, business_id, employee_id)
    if employee is None:
        raise LookupError("employee_not_found")
    return employee


async def _validate_roles(
    session: AsyncSession,
    business_id: UUID,
    role_ids: list[UUID],
) -> dict[UUID, Role]:
    if len(role_ids) != len(set(role_ids)):
        raise ValueError("duplicate_role_ids")
    if not role_ids:
        return {}

    result = await session.execute(
        select(Role).where(Role.business_id == business_id, Role.id.in_(role_ids))
    )
    roles = {role.id: role for role in result.scalars().all()}
    if len(roles) != len(role_ids):
        raise LookupError("role_not_found")
    return roles


async def _validate_locations(
    session: AsyncSession,
    business_id: UUID,
    location_ids: list[UUID],
) -> dict[UUID, Location]:
    if len(location_ids) != len(set(location_ids)):
        raise ValueError("duplicate_location_ids")
    if not location_ids:
        return {}

    result = await session.execute(
        select(Location).where(Location.business_id == business_id, Location.id.in_(location_ids))
    )
    locations = {location.id: location for location in result.scalars().all()}
    if len(locations) != len(location_ids):
        raise LookupError("location_not_found")
    return locations


def _primary_role_id(assignments: list[EmployeeRoleUpsert]) -> UUID | None:
    if not assignments:
        return None
    primary_ids = [item.role_id for item in assignments if item.is_primary]
    if len(primary_ids) > 1:
        raise ValueError("multiple_primary_roles")
    if primary_ids:
        return primary_ids[0]
    return assignments[0].role_id


def _primary_location_id(assignments: list[EmployeeLocationUpsert]) -> UUID | None:
    if not assignments:
        return None
    primary_ids = [item.location_id for item in assignments if item.is_primary]
    if len(primary_ids) > 1:
        raise ValueError("multiple_primary_locations")
    if primary_ids:
        return primary_ids[0]
    return assignments[0].location_id


async def _replace_employee_roles(
    session: AsyncSession,
    employee: Employee,
    business_id: UUID,
    assignments: list[EmployeeRoleUpsert],
) -> None:
    await _validate_roles(session, business_id, [item.role_id for item in assignments])
    primary_role_id = _primary_role_id(assignments)

    existing_roles = await _list_employee_roles(session, employee.id)
    existing_by_role_id = {item.role_id: item for item in existing_roles}
    desired_role_ids = {item.role_id for item in assignments}

    for existing in existing_roles:
        if existing.role_id not in desired_role_ids:
            await session.delete(existing)
            continue
        existing.is_primary = False

    await session.flush()

    for item in assignments:
        existing = existing_by_role_id.get(item.role_id)
        if existing is None:
            record = EmployeeRole(
                employee_id=employee.id,
                role_id=item.role_id,
                proficiency_level=item.proficiency_level or 1,
                is_primary=item.role_id == primary_role_id,
                role_metadata=item.role_metadata or {},
            )
            session.add(record)
            continue

        existing.is_primary = existing.role_id == primary_role_id
        if "proficiency_level" in item.model_fields_set and item.proficiency_level is not None:
            existing.proficiency_level = item.proficiency_level
        if "role_metadata" in item.model_fields_set:
            existing.role_metadata = item.role_metadata or {}


async def _replace_employee_locations(
    session: AsyncSession,
    employee: Employee,
    business_id: UUID,
    assignments: list[EmployeeLocationUpsert],
) -> None:
    await _validate_locations(session, business_id, [item.location_id for item in assignments])
    primary_location_id = _primary_location_id(assignments)

    existing_locations = await _list_employee_locations(session, employee.id)
    existing_by_location_id = {item.location_id: item for item in existing_locations}
    desired_location_ids = {item.location_id for item in assignments}

    for existing in existing_locations:
        if existing.location_id not in desired_location_ids:
            await session.delete(existing)
            continue
        existing.is_primary = False

    await session.flush()

    for item in assignments:
        existing = existing_by_location_id.get(item.location_id)
        if existing is None:
            record = EmployeeLocation(
                employee_id=employee.id,
                location_id=item.location_id,
                is_primary=item.location_id == primary_location_id,
                access_level=item.access_level or "approved",
                location_source=item.location_source,
                can_cover_last_minute=(
                    True if item.can_cover_last_minute is None else item.can_cover_last_minute
                ),
                can_blast=True if item.can_blast is None else item.can_blast,
                travel_radius_miles=item.travel_radius_miles,
                location_metadata=item.location_metadata or {},
            )
            session.add(record)
            continue

        existing.is_primary = existing.location_id == primary_location_id
        if "access_level" in item.model_fields_set and item.access_level is not None:
            existing.access_level = item.access_level
        if "location_source" in item.model_fields_set:
            existing.location_source = item.location_source
        if "can_cover_last_minute" in item.model_fields_set and item.can_cover_last_minute is not None:
            existing.can_cover_last_minute = item.can_cover_last_minute
        if "can_blast" in item.model_fields_set and item.can_blast is not None:
            existing.can_blast = item.can_blast
        if "travel_radius_miles" in item.model_fields_set:
            existing.travel_radius_miles = item.travel_radius_miles
        if "location_metadata" in item.model_fields_set:
            existing.location_metadata = item.location_metadata or {}


async def _replace_employee_work_permits(
    session: AsyncSession,
    employee: Employee,
    permits: list[EmployeeWorkPermitCreate],
) -> list[EmployeeWorkPermit]:
    normalized_permit_numbers = [str(permit.permit_number or "").strip() for permit in permits]
    if any(not permit_number for permit_number in normalized_permit_numbers):
        raise ValueError("invalid_work_permit_number")
    if len(set(normalized_permit_numbers)) != len(normalized_permit_numbers):
        raise ValueError("duplicate_work_permit_numbers")

    _hydrate_employee_defaults(employee)
    try:
        existing_work_permits = await _list_employee_work_permits(session, employee.id)
    except AttributeError:
        existing_work_permits = list(getattr(employee, "work_permits", []) or [])
    for existing in existing_work_permits:
        await session.delete(existing)
    await session.flush()

    records: list[EmployeeWorkPermit] = []
    for permit, normalized_permit_number in zip(permits, normalized_permit_numbers):
        permit_metadata = dict(permit.permit_metadata or {})
        permit_metadata.pop("rule_profile", None)
        rule_profile = _serialized_work_permit_rule_profile(permit.rule_profile)
        if rule_profile is not None:
            permit_metadata["rule_profile"] = rule_profile
        record = EmployeeWorkPermit(
            employee_id=employee.id,
            permit_number=normalized_permit_number,
            issuing_authority=str(permit.issuing_authority or "").strip() or None,
            issued_on=permit.issued_on,
            effective_start_date=permit.effective_start_date,
            effective_end_date=permit.effective_end_date,
            max_daily_minutes=permit.max_daily_minutes,
            max_weekly_minutes=permit.max_weekly_minutes,
            earliest_start_local_time=permit.earliest_start_local_time,
            latest_end_local_time=permit.latest_end_local_time,
            permit_metadata=permit_metadata,
        )
        _hydrate_work_permit_defaults(employee, record)
        session.add(record)
        records.append(record)
    await session.flush()
    _sync_employee_work_permit_snapshot(employee, records)
    return records


async def create_employee(
    session: AsyncSession,
    business_id: UUID,
    payload: EmployeeCreate,
    *,
    linked_user_id: UUID | None = None,
) -> Employee:
    business = await _require_business(session, business_id)
    normalized_external_ref = _normalize_employee_external_ref(payload.external_ref)
    normalized_phone = _normalize_employee_phone(payload.phone_e164)
    normalized_email = _normalize_employee_email(payload.email)
    duplicate = await _find_duplicate_employee(
        session,
        business_id,
        external_ref=normalized_external_ref,
        phone_e164=normalized_phone,
        email=normalized_email,
    )
    if duplicate is not None:
        duplicate_field, _ = duplicate
        raise ValueError(f"employee_duplicate_{duplicate_field}")
    primary_business_location: Location | None = None
    if payload.primary_location_id is not None:
        primary_business_location = await session.get(Location, payload.primary_location_id)
        if primary_business_location is None or primary_business_location.business_id != business_id:
            raise LookupError("primary_location_not_found")

    employee = Employee(
        business_id=business_id,
        user_id=linked_user_id,
        external_ref=normalized_external_ref,
        employee_number=payload.employee_number,
        full_name=payload.full_name,
        preferred_name=payload.preferred_name,
        phone_e164=normalized_phone,
        email=normalized_email,
        base_hourly_rate_cents=payload.base_hourly_rate_cents,
        compliance_regular_rate_cents=payload.compliance_regular_rate_cents,
        date_of_birth=payload.date_of_birth,
        minor_school_status=_normalize_minor_school_status(payload.minor_school_status),
        work_permit_number=str(payload.work_permit_number or "").strip() or None,
        work_permit_effective_start_on=payload.work_permit_effective_start_on,
        work_permit_expires_on=payload.work_permit_expires_on,
        work_permit_max_daily_minutes=payload.work_permit_max_daily_minutes,
        work_permit_max_weekly_minutes=payload.work_permit_max_weekly_minutes,
        work_permit_earliest_start_local_time=payload.work_permit_earliest_start_local_time,
        work_permit_latest_end_local_time=payload.work_permit_latest_end_local_time,
        employment_type=payload.employment_type,
        hire_date=payload.hire_date,
        notes=payload.notes,
        employee_metadata=payload.employee_metadata,
    )
    _hydrate_employee_defaults(employee)
    session.add(employee)
    await session.flush()
    primary_location: EmployeeLocation | None = None
    if payload.primary_location_id is not None:
        primary_location = EmployeeLocation(
            employee_id=employee.id,
            location_id=payload.primary_location_id,
            is_primary=True,
            access_level="approved",
            location_source="employee_create",
            can_cover_last_minute=True,
            can_blast=True,
            location_metadata={},
        )
        session.add(primary_location)

    availability_timezone = _default_availability_timezone(
        business=business,
        primary_location=primary_business_location,
    )
    default_availability_rules = _default_weekly_availability_rules(
        employee_id=employee.id,
        timezone_name=availability_timezone,
        source="employee_create",
    )
    for rule in default_availability_rules:
        session.add(rule)

    work_permits: list[EmployeeWorkPermit] = []
    if payload.work_permits:
        work_permits = await _replace_employee_work_permits(session, employee, payload.work_permits)

    await session.flush()
    await session.refresh(employee)
    if primary_location is not None:
        if primary_business_location is not None:
            set_committed_value(primary_location, "location", primary_business_location)
        set_committed_value(employee, "employee_locations", [primary_location])
    else:
        set_committed_value(employee, "employee_locations", [])
    set_committed_value(employee, "employee_roles", [])
    set_committed_value(employee, "work_permits", work_permits)
    set_committed_value(employee, "availability_rules", default_availability_rules)
    return _attach_employee_notification_preferences(employee)


async def bulk_import_employees(
    session: AsyncSession,
    business_id: UUID,
    *,
    filename: str,
    content: bytes,
) -> EmployeeBulkImportRead:
    await _require_business(session, business_id)

    if not content:
        raise ValueError("employee_import_file_empty")

    location_result = await session.execute(
        select(Location)
        .where(Location.business_id == business_id, Location.is_active.is_(True))
        .order_by(Location.created_at.asc())
    )
    active_locations = list(location_result.scalars().all())
    default_location = active_locations[0] if len(active_locations) == 1 else None

    parsed_rows, errors = parse_employee_import_file(
        filename,
        content,
        default_location_id=default_location.id if default_location is not None else None,
    )

    created_employees: list[Employee] = []
    for row in parsed_rows:
        try:
            created_employees.append(await create_employee(session, business_id, row))
        except ValueError as exc:
            detail = str(exc)
            if detail.startswith("employee_duplicate_"):
                source_row_number = (row.employee_metadata or {}).get("source_row_number")
                errors.append(
                    EmployeeImportErrorRead(
                        row_number=source_row_number if isinstance(source_row_number, int) else None,
                        message=detail,
                    )
                )
                continue
            raise

    return EmployeeBulkImportRead(
        created_count=len(created_employees),
        skipped_count=len(errors),
        employees=created_employees,
        errors=errors,
        default_location_id=default_location.id if default_location is not None else None,
        default_location_name=default_location.location_display_name if default_location is not None else None,
    )


async def enroll_employee_at_location(
    session: AsyncSession,
    business_id: UUID,
    payload: EmployeeEnrollAtLocationCreate,
) -> EmployeeEnrollmentRead:
    await _require_business(session, business_id)

    location = await session.get(Location, payload.location_id)
    if location is None or location.business_id != business_id:
        raise LookupError("location_not_found")
    if not payload.role_ids:
        raise ValueError("role_ids_required")

    roles = []
    for role_id in payload.role_ids:
        role = await session.get(Role, role_id)
        if role is None or role.business_id != business_id:
            raise LookupError("role_not_found")
        roles.append(role)

    employee = Employee(
        business_id=business_id,
        external_ref=payload.external_ref,
        employee_number=payload.employee_number,
        full_name=payload.full_name,
        preferred_name=payload.preferred_name,
        phone_e164=payload.phone_e164,
        email=payload.email,
        base_hourly_rate_cents=payload.base_hourly_rate_cents,
        compliance_regular_rate_cents=payload.compliance_regular_rate_cents,
        date_of_birth=payload.date_of_birth,
        minor_school_status=_normalize_minor_school_status(payload.minor_school_status),
        work_permit_number=str(payload.work_permit_number or "").strip() or None,
        work_permit_effective_start_on=payload.work_permit_effective_start_on,
        work_permit_expires_on=payload.work_permit_expires_on,
        work_permit_max_daily_minutes=payload.work_permit_max_daily_minutes,
        work_permit_max_weekly_minutes=payload.work_permit_max_weekly_minutes,
        work_permit_earliest_start_local_time=payload.work_permit_earliest_start_local_time,
        work_permit_latest_end_local_time=payload.work_permit_latest_end_local_time,
        employment_type=payload.employment_type,
        hire_date=payload.hire_date,
        notes=payload.notes,
        employee_metadata=payload.employee_metadata,
    )
    _hydrate_employee_defaults(employee)
    session.add(employee)
    await session.flush()

    primary_employee_location = EmployeeLocation(
        employee_id=employee.id,
        location_id=payload.location_id,
        is_primary=True,
        access_level="approved",
        location_source="location_enrollment",
        can_cover_last_minute=True,
        can_blast=True,
        location_metadata={},
    )
    session.add(primary_employee_location)

    employee_roles: list[EmployeeRole] = []
    for index, role in enumerate(roles):
        employee_role = EmployeeRole(
            employee_id=employee.id,
            role_id=role.id,
            proficiency_level=1,
            is_primary=index == 0,
            role_metadata={"source": "location_enrollment"},
        )
        session.add(employee_role)
        employee_roles.append(employee_role)

    work_permits: list[EmployeeWorkPermit] = []
    if payload.work_permits:
        work_permits = await _replace_employee_work_permits(session, employee, payload.work_permits)

    await session.flush()
    await session.refresh(employee)
    set_committed_value(primary_employee_location, "location", location)
    for employee_role, role in zip(employee_roles, roles):
        set_committed_value(employee_role, "role", role)
    set_committed_value(employee, "employee_locations", [primary_employee_location])
    set_committed_value(employee, "employee_roles", employee_roles)
    set_committed_value(employee, "work_permits", work_permits)

    return EmployeeEnrollmentRead(
        employee=_attach_employee_notification_preferences(employee),
        roles=employee_roles,
    )


async def get_employee_profile(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
) -> Employee:
    employee = await _require_employee(session, business_id, employee_id)
    employee_roles = await _list_employee_roles(session, employee_id)
    employee_locations = await _list_employee_locations(session, employee_id)
    try:
        employee_work_permits = await _list_employee_work_permits(session, employee_id)
    except AttributeError:
        employee_work_permits = list(getattr(employee, "work_permits", []) or [])
    set_committed_value(employee, "employee_roles", employee_roles)
    set_committed_value(employee, "employee_locations", employee_locations)
    set_committed_value(employee, "work_permits", employee_work_permits)
    return _attach_employee_notification_preferences(employee)


async def _blocking_employee_assignment_count(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
) -> int:
    count = await session.scalar(
        select(func.count(ShiftAssignment.id))
        .select_from(ShiftAssignment)
        .join(Shift, Shift.id == ShiftAssignment.shift_id)
        .where(
            Shift.business_id == business_id,
            ShiftAssignment.employee_id == employee_id,
            Shift.lifecycle_status.not_in(
                [ShiftLifecycleStatus.draft, ShiftLifecycleStatus.cancelled]
            ),
            ShiftAssignment.status.in_(EMPLOYEE_DELETE_BLOCKING_ASSIGNMENT_STATUSES),
        )
    )
    return int(count or 0)


async def get_employee_delete_readiness(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
) -> dict[str, object]:
    await _require_employee(session, business_id, employee_id)
    blocking_assignment_count = await _blocking_employee_assignment_count(
        session,
        business_id,
        employee_id,
    )
    reason = (
        "This employee has scheduled shifts and cannot be removed until those shifts are cleared."
        if blocking_assignment_count > 0
        else None
    )
    return {
        "business_id": business_id,
        "employee_id": employee_id,
        "can_delete": blocking_assignment_count == 0,
        "reason": reason,
    }


async def delete_employee(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
) -> Employee:
    employee = await _require_employee(session, business_id, employee_id)
    blocking_assignment_count = await _blocking_employee_assignment_count(
        session,
        business_id,
        employee_id,
    )
    if blocking_assignment_count > 0:
        raise ValueError("employee_has_operational_data")
    await session.delete(employee)
    await session.flush()
    return employee


async def update_employee(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
    payload: EmployeeUpdate,
) -> Employee:
    employee = await _require_employee(session, business_id, employee_id)
    normalized_external_ref = (
        _normalize_employee_external_ref(payload.external_ref)
        if "external_ref" in payload.model_fields_set
        else employee.external_ref
    )
    normalized_phone = (
        _normalize_employee_phone(payload.phone_e164)
        if "phone_e164" in payload.model_fields_set
        else employee.phone_e164
    )
    normalized_email = (
        _normalize_employee_email(payload.email)
        if "email" in payload.model_fields_set
        else employee.email
    )

    if {
        "external_ref",
        "phone_e164",
        "email",
    }.intersection(payload.model_fields_set):
        duplicate = await _find_duplicate_employee(
            session,
            business_id,
            external_ref=normalized_external_ref,
            phone_e164=normalized_phone,
            email=normalized_email,
            exclude_employee_id=employee_id,
        )
        if duplicate is not None:
            duplicate_field, _ = duplicate
            raise ValueError(f"employee_duplicate_{duplicate_field}")

    field_updates = {
        "full_name": payload.full_name,
        "preferred_name": payload.preferred_name,
        "phone_e164": normalized_phone,
        "email": normalized_email,
        "external_ref": normalized_external_ref,
        "employee_number": payload.employee_number,
        "base_hourly_rate_cents": payload.base_hourly_rate_cents,
        "compliance_regular_rate_cents": payload.compliance_regular_rate_cents,
        "date_of_birth": payload.date_of_birth,
        "minor_school_status": (
            _normalize_minor_school_status(payload.minor_school_status)
            if "minor_school_status" in payload.model_fields_set
            else employee.minor_school_status
        ),
        "work_permit_number": (
            str(payload.work_permit_number or "").strip() or None
            if "work_permit_number" in payload.model_fields_set
            else employee.work_permit_number
        ),
        "work_permit_effective_start_on": payload.work_permit_effective_start_on,
        "work_permit_expires_on": payload.work_permit_expires_on,
        "work_permit_max_daily_minutes": payload.work_permit_max_daily_minutes,
        "work_permit_max_weekly_minutes": payload.work_permit_max_weekly_minutes,
        "work_permit_earliest_start_local_time": payload.work_permit_earliest_start_local_time,
        "work_permit_latest_end_local_time": payload.work_permit_latest_end_local_time,
        "employment_type": payload.employment_type,
        "status": payload.status,
        "hire_date": payload.hire_date,
        "termination_date": payload.termination_date,
        "notes": payload.notes,
        "employee_metadata": payload.employee_metadata,
    }
    for field_name, field_value in field_updates.items():
        if field_name in payload.model_fields_set:
            setattr(employee, field_name, field_value)

    if "notification_preferences" in payload.model_fields_set and payload.notification_preferences is not None:
        _apply_employee_notification_preferences_update(employee, payload.notification_preferences)

    if "roles" in payload.model_fields_set:
        await _replace_employee_roles(session, employee, business_id, payload.roles or [])

    if "locations" in payload.model_fields_set:
        await _replace_employee_locations(session, employee, business_id, payload.locations or [])

    if "work_permits" in payload.model_fields_set and payload.work_permits is not None:
        await _replace_employee_work_permits(session, employee, payload.work_permits)

    await session.flush()
    refreshed = await get_employee(session, business_id, employee_id)
    if refreshed is None:
        raise LookupError("employee_not_found")
    employee_roles = await _list_employee_roles(session, employee_id)
    employee_locations = await _list_employee_locations(session, employee_id)
    try:
        employee_work_permits = await _list_employee_work_permits(session, employee_id)
    except AttributeError:
        employee_work_permits = list(getattr(refreshed, "work_permits", []) or [])
    set_committed_value(refreshed, "employee_roles", employee_roles)
    set_committed_value(refreshed, "employee_locations", employee_locations)
    set_committed_value(refreshed, "work_permits", employee_work_permits)
    return _attach_employee_notification_preferences(refreshed)


async def add_employee_role(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
    payload: EmployeeRoleCreate,
) -> EmployeeRole:
    employee = await session.get(Employee, employee_id)
    role = await session.get(Role, payload.role_id)
    if employee is None or employee.business_id != business_id or role is None or role.business_id != business_id:
        raise LookupError("employee_or_role_not_found")

    existing_roles = await _list_employee_roles(session, employee_id)
    existing_role = next((item for item in existing_roles if item.role_id == payload.role_id), None)
    if payload.is_primary:
        for employee_role in existing_roles:
            if employee_role.role_id != payload.role_id and employee_role.is_primary:
                employee_role.is_primary = False
        await session.flush()

    if existing_role is not None:
        existing_role.proficiency_level = payload.proficiency_level
        existing_role.is_primary = (
            payload.is_primary
            or existing_role.is_primary
            or not any(item.is_primary for item in existing_roles)
        )
        existing_role.role_metadata = payload.role_metadata
        await session.flush()
        await session.refresh(existing_role)
        return existing_role

    record = EmployeeRole(
        employee_id=employee_id,
        role_id=payload.role_id,
        proficiency_level=payload.proficiency_level,
        is_primary=payload.is_primary or not any(item.is_primary for item in existing_roles),
        role_metadata=payload.role_metadata,
    )
    session.add(record)
    await session.flush()
    await session.refresh(record)
    return record


async def add_employee_location(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
    payload: EmployeeLocationCreate,
) -> EmployeeLocation:
    employee = await session.get(Employee, employee_id)
    location = await session.get(Location, payload.location_id)
    if employee is None or employee.business_id != business_id or location is None or location.business_id != business_id:
        raise LookupError("employee_or_location_not_found")

    existing_locations = await _list_employee_locations(session, employee_id)
    existing_location = next(
        (
            item
            for item in existing_locations
            if item.location_id == payload.location_id
        ),
        None,
    )
    has_primary_location = any(item.is_primary for item in existing_locations)

    if payload.is_primary:
        for employee_location in existing_locations:
            if employee_location.location_id != payload.location_id and employee_location.is_primary:
                employee_location.is_primary = False
        await session.flush()

    if existing_location is not None:
        existing_location.is_primary = (
            existing_location.is_primary
            or payload.is_primary
            or not has_primary_location
        )
        existing_location.access_level = payload.access_level
        existing_location.location_source = payload.location_source
        existing_location.can_cover_last_minute = payload.can_cover_last_minute
        existing_location.can_blast = payload.can_blast
        existing_location.travel_radius_miles = payload.travel_radius_miles
        existing_location.location_metadata = payload.location_metadata
        await session.flush()
        await session.refresh(existing_location)
        return existing_location

    record = EmployeeLocation(
        employee_id=employee_id,
        location_id=payload.location_id,
        is_primary=payload.is_primary or not has_primary_location,
        access_level=payload.access_level,
        location_source=payload.location_source,
        can_cover_last_minute=payload.can_cover_last_minute,
        can_blast=payload.can_blast,
        travel_radius_miles=payload.travel_radius_miles,
        location_metadata=payload.location_metadata,
    )
    session.add(record)
    await session.flush()
    await session.refresh(record)
    return record


async def add_employee_availability_rule(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
    payload: EmployeeAvailabilityRuleCreate,
) -> EmployeeAvailabilityRule:
    employee = await session.get(Employee, employee_id)
    if employee is None or employee.business_id != business_id:
        raise LookupError("employee_not_found")

    record = EmployeeAvailabilityRule(
        employee_id=employee_id,
        day_of_week=payload.day_of_week,
        start_local_time=payload.start_local_time,
        end_local_time=payload.end_local_time,
        timezone=payload.timezone,
        availability_type=payload.availability_type,
        valid_from=payload.valid_from,
        valid_until=payload.valid_until,
        priority=payload.priority,
        availability_metadata=payload.availability_metadata,
    )
    session.add(record)
    await session.flush()
    await session.refresh(record)
    return record


async def replace_employee_availability_rules(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
    payload: EmployeeAvailabilityRuleReplace,
) -> list[EmployeeAvailabilityRule]:
    employee = await session.get(Employee, employee_id)
    if employee is None or employee.business_id != business_id:
        raise LookupError("employee_not_found")

    await session.execute(
        delete(EmployeeAvailabilityRule).where(
            EmployeeAvailabilityRule.employee_id == employee_id,
        )
    )

    for rule in payload.rules:
        session.add(
            EmployeeAvailabilityRule(
                employee_id=employee_id,
                day_of_week=rule.day_of_week,
                start_local_time=rule.start_local_time,
                end_local_time=rule.end_local_time,
                timezone=rule.timezone,
                availability_type=rule.availability_type,
                valid_from=rule.valid_from,
                valid_until=rule.valid_until,
                priority=rule.priority,
                availability_metadata=rule.availability_metadata,
            )
        )

    await session.flush()
    return await _list_employee_availability_rules(session, employee_id)


async def get_employee_availability_rules(
    session: AsyncSession,
    business_id: UUID,
    employee_id: UUID,
) -> tuple[Employee, list[EmployeeAvailabilityRule]]:
    employee = await session.get(Employee, employee_id)
    if employee is None or employee.business_id != business_id:
        raise LookupError("employee_not_found")
    rules = await _list_employee_availability_rules(session, employee_id)
    return employee, rules


async def get_self_employee_availability_rules(
    session: AsyncSession,
    business_id: UUID,
    *,
    user_id: UUID | None,
    email: str | None,
    phone_e164: str | None,
    full_name: str | None,
) -> tuple[Employee, list[EmployeeAvailabilityRule]]:
    employee = await _resolve_self_service_employee(
        session,
        business_id,
        user_id=user_id,
        email=email,
        phone_e164=phone_e164,
        full_name=full_name,
    )
    rules = await _list_employee_availability_rules(session, employee.id)
    return employee, rules


async def replace_self_employee_availability_rules(
    session: AsyncSession,
    business_id: UUID,
    *,
    user_id: UUID | None,
    email: str | None,
    phone_e164: str | None,
    full_name: str | None,
    payload: EmployeeAvailabilityRuleReplace,
) -> tuple[Employee, list[EmployeeAvailabilityRule]]:
    employee = await _get_or_create_self_service_employee(
        session,
        business_id,
        user_id=user_id,
        email=email,
        phone_e164=phone_e164,
        full_name=full_name,
    )
    rules = await replace_employee_availability_rules(
        session,
        business_id,
        employee.id,
        payload,
    )
    return employee, rules


from app.models.scheduling import Shift, ShiftAssignment  # noqa: E402
